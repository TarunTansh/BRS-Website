from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
import re
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT_DIR = Path('xls')
OUTPUT = INPUT_DIR / 'BRS_merged_job_fair_data.xlsx'
TODAY = date.today()
ALIASES = {
    'location': ['HOMETOWN', 'LOCATION', 'Home town', 'Location'],
    'candidate_name': ['CANDIDATE NAME', 'Candidate Name', 'Name of Employees', 'NAME'],
    'phone': ['CONTACT NUMBER', 'Contact no.', 'Mobile No', 'Contact no', 'Contact .No', 'CONTACT NO', 'Phone Number'],
    'qualification': ['QUALIFICATION', 'Qualitification', 'Qualification', 'Highest Qualification'],
    'gender': ['Gender'],
    'date_of_birth': ['D.O.B', 'Date of Birth'],
    'experience_years': ['Experience Year'],
    'experience_months': ['Experience Month'],
    'client_or_job': ['CLIENT NAME', 'Shortlisted For (Client Name)', 'Client to be hired on', 'Lineup For', 'Hired Designation'],
    'remarks': ['REMARKS', 'REMARK', 'Remarks'],
    'rc_name': ['RC NAME', 'INTERN', 'DEEPAK'],
}
FINAL_HEADERS = ['Candidate Name', 'Phone Number', 'Location', 'Highest Qualification', 'Gender', 'Date of Birth', 'Age', 'Experience Years', 'Experience Months', 'Client Name', 'Remarks', 'RC Name']
CANONICAL = ['source_file', 'source_sheet', 'source_row', 'location', 'candidate_name', 'phone', 'qualification', 'gender', 'date_of_birth', 'age', 'experience_years', 'experience_months', 'client_or_job', 'remarks', 'rc_name', 'duplicate_group_size']

def norm(value):
    return re.sub(r'\s+', ' ', str(value).strip().lower()) if value is not None else ''

def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()

def clean_phone(value):
    digits = re.sub(r'\D', '', text(value))
    return digits[-10:] if len(digits) >= 10 else digits

def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%d.%m.%Y', '%d/%m/%y', '%d-%m-%y'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None

def calculate_age(value):
    dob = parse_date(value)
    if not dob or dob > TODAY:
        return ''
    return TODAY.year - dob.year - ((TODAY.month, TODAY.day) < (dob.month, dob.day))

def mapping(headers):
    normalized = {norm(header): index for index, header in enumerate(headers)}
    return {field: next((normalized[norm(alias)] for alias in aliases if norm(alias) in normalized), None) for field, aliases in ALIASES.items()}

def style(sheet):
    fill = PatternFill('solid', fgColor='128C7E')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30
    for cells in sheet.columns:
        letter = get_column_letter(cells[0].column)
        width = min(max(max((len(str(cell.value or '')) for cell in cells[:100]), default=10) + 2, 12), 36)
        sheet.column_dimensions[letter].width = width

def read_records():
    records, summaries = [], []
    for path in sorted(INPUT_DIR.glob('*.xlsx')):
        if path.name == OUTPUT.name or path.name.startswith('~$'):
            continue
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in book.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            maps = mapping(headers)
            count = 0
            for source_row, values in enumerate(rows, 2):
                values = list(values)
                if not any(text(value) for value in values):
                    continue
                record = {field: text(values[index]) if index is not None and index < len(values) else '' for field, index in maps.items()}
                record['phone'] = clean_phone(record['phone'])
                record['age'] = calculate_age(record['date_of_birth'])
                record.update(source_file=path.name, source_sheet=sheet.title, source_row=source_row)
                record['extras'] = {str(header).strip() if header is not None and str(header).strip() else f'column_{index + 1}': text(values[index]) if index < len(values) else '' for index, header in enumerate(headers)}
                records.append(record)
                count += 1
            summaries.append([path.name, sheet.title, count, ', '.join(text(header) for header in headers if header is not None)])
    return records, summaries

def build(records, summaries):
    groups = OrderedDict()
    for record in records:
        key = record['phone'] or f"{record['source_file']}:{record['source_sheet']}:{record['source_row']}"
        groups.setdefault(key, []).append(record)
    book = Workbook()
    readme = book.active
    readme.title = 'Read Me'
    for row in [
        ['BRS merged job fair data'],
        ['Final columns', ', '.join(FINAL_HEADERS)],
        ['Source records', len(records)],
        ['Unified candidates', len(groups)],
        ['Age calculation', f'Calculated from Date of Birth as of {TODAY.isoformat()}.'],
        ['RC Name rule', 'Blank unless an actual RC NAME, INTERN, or DEEPAK value exists in the source row.'],
    ]:
        readme.append(row)
    readme.merge_cells('A1:B1')
    readme['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    readme['A1'].fill = PatternFill('solid', fgColor='128C7E')
    readme.column_dimensions['A'].width = 25
    readme.column_dimensions['B'].width = 110
    for row in readme.iter_rows(min_row=2): row[0].font = Font(bold=True)

    unified = book.create_sheet('Unified Candidates')
    unified.append(FINAL_HEADERS)
    for group in groups.values():
        merged = {field: next((item[field] for item in group if item[field] != ''), '') for field in ALIASES}
        merged['age'] = calculate_age(merged['date_of_birth'])
        unified.append([
            merged['candidate_name'], merged['phone'], merged['location'], merged['qualification'], merged['gender'], merged['date_of_birth'], merged['age'], merged['experience_years'], merged['experience_months'], merged['client_or_job'], merged['remarks'], merged['rc_name']
        ])
    style(unified)

    all_sheet = book.create_sheet('All Source Records')
    all_sheet.append(FINAL_HEADERS + CANONICAL)
    for record in records:
        group_size = len(groups.get(record['phone'], [record])) if record['phone'] else 1
        all_sheet.append([
            record['candidate_name'], record['phone'], record['location'], record['qualification'], record['gender'], record['date_of_birth'], record['age'], record['experience_years'], record['experience_months'], record['client_or_job'], record['remarks'], record['rc_name'], record['source_file'], record['source_sheet'], record['source_row'], record['location'], record['candidate_name'], record['phone'], record['qualification'], record['gender'], record['date_of_birth'], record['age'], record['experience_years'], record['experience_months'], record['client_or_job'], record['remarks'], record['rc_name'], group_size
        ])
    style(all_sheet)

    summary = book.create_sheet('Source Summary')
    summary.append(['source_file', 'source_sheet', 'records', 'headers'])
    for row in summaries: summary.append(row)
    style(summary)
    book.save(OUTPUT)
    print(f'Updated {OUTPUT}')
    print(f'Source records: {len(records)}')
    print(f'Unified candidates: {len(groups)}')
    print(f'RC Name populated: {sum(bool(next((item["rc_name"] for item in group if item["rc_name"]), "")) for group in groups.values())}')

records, summaries = read_records()
build(records, summaries)
